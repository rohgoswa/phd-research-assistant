import streamlit as st
import os
import platform
from docx import Document
from io import BytesIO
from langchain_groq import ChatGroq
from langchain_community.chat_models import ChatOllama
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import HuggingFaceEmbeddings
from pypdf import PdfReader

# --- PAGE CONFIG ---
st.set_page_config(page_title="Sarkari Decoder", page_icon="🇮🇳", layout="wide")

# --- CUSTOM CSS ---
st.markdown("""
<style>
    .stSpinner {text-align: center; color: #4CAF50;}
    div[data-testid="stSidebar"] {background-color: #f7f9fb;}
    [data-testid="stSidebar"] img {
        border-radius: 50%;
        object-fit: cover;
    }
</style>
""", unsafe_allow_html=True)

# --- 1. INITIALIZE SESSION STATE ---
if "messages" not in st.session_state:
    st.session_state.messages = []
if "vector_store" not in st.session_state:
    st.session_state.vector_store = None

# --- 2. HELPER FUNCTIONS ---
def create_docx(messages):
    doc = Document()
    doc.add_heading('Research Session Report', 0)
    for msg in messages:
        role = "User" if msg["role"] == "user" else "AI Assistant"
        doc.add_heading(role, level=2)
        doc.add_paragraph(msg["content"])
        doc.add_paragraph("-" * 20)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# --- 3. SIDEBAR CONFIG ---
with st.sidebar:
    col1, col2, col3 = st.columns([1,2,1])
    with col2:
        if os.path.exists("profile.png"):
            st.image("profile.png", width=100) 
    
    st.title("Settings")
    st.write("Built by **Rohit Goswami**") 
    
    # Mode Selection
    os_name = platform.system()
    mode_options = ["☁️ Cloud (Speed)"]
    if os_name == "Darwin":
        mode_options.insert(0, "🔒 Local (Privacy)")
    mode = st.radio("Processing Mode:", mode_options)
    
    if mode == "☁️ Cloud (Speed)":
        api_key = st.text_input("Groq API Key", type="password", placeholder="gsk_...")
        st.caption("[Get Key](https://console.groq.com/keys)")
        model_id = "llama-3.3-70b-versatile" 
    else:
        model_id = st.selectbox("Local Model:", ["llama3.2", "deepseek-r1"], index=0)
        st.info("Mode: 100% Private (Runs on your Mac)")
    
    st.divider()
    
    # BUTTON 1: Clear Memory
    if st.button("🗑️ Clear Chat History"):
        st.session_state.messages = []
        st.session_state.vector_store = None
        st.rerun()

    # BUTTON 2: Download Report
    if st.session_state.messages:
        docx_file = create_docx(st.session_state.messages)
        st.download_button(
            label="📥 Download Report",
            data=docx_file,
            file_name="research_session.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

# --- 4. APP LOGIC ---
st.title("🇮🇳 Sarkari Decoder") 
st.caption("Upload confusing Govt. Notices, Circulars & Gazettes. Get simple explanations in seconds.")

# UPDATED: Accept Multiple Files
uploaded_files = st.file_uploader(
    "Upload PDF(s)", 
    type="pdf", 
    label_visibility="collapsed", 
    accept_multiple_files=True 
)

# Process files loop
if uploaded_files and st.session_state.vector_store is None:
    with st.spinner(f"🧠 Reading {len(uploaded_files)} documents..."):
        all_text = ""
        for uploaded_file in uploaded_files:
            pdf_reader = PdfReader(uploaded_file)
            for page in pdf_reader.pages:
                all_text += page.extract_text()
            
        # Updated Text Splitter with Overlap
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1200, chunk_overlap=500, length_function=len)
        chunks = text_splitter.split_text(all_text)
        
        embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")
        st.session_state.vector_store = FAISS.from_texts(chunks, embeddings)
        st.toast(f"Indexed {len(uploaded_files)} Documents Successfully!", icon="✅")

# Display Chat
for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

# Handle Input
if st.session_state.vector_store:
    if prompt := st.chat_input("Ask a question about your documents..."):
        st.session_state.messages.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)

        with st.chat_message("assistant"):
            with st.spinner("Thinking..."):
                # Updated k=10 for deeper search
                docs = st.session_state.vector_store.similarity_search(prompt, k=10)
                context = "\n\n".join([doc.page_content for doc in docs])
                
                # Strict Prompt
                rag_prompt = f"""You are a helpful Government Document Simplifier. 
                Your job is to explain the confusing government circular/rule to the user in SIMPLE terms.
                
                RULES:
                1. Use simple, easy-to-understand English.
                2. If there are dates or deadlines, list them clearly.
                3. If there are fees, bold them.
                4. Do not use complex jargon.
                
                CONTEXT:
                {context}
                
                QUESTION: 
                {prompt}
                """
                
                try:
                    if mode == "☁️ Cloud (Speed)":
                        if not api_key:
                            st.error("Please enter API Key in sidebar!")
                            st.stop()
                        llm = ChatGroq(groq_api_key=api_key, model_name=model_id)
                        response = llm.invoke(rag_prompt)
                        result = response.content
                    else:
                        llm = ChatOllama(model=model_id)
                        response = llm.invoke(rag_prompt)
                        result = response.content
                    
                    st.markdown(result)
                    st.session_state.messages.append({"role": "assistant", "content": result})
                    
                    with st.expander("📚 View Sources"):
                        st.caption(context)
                        
                except Exception as e:
                    st.error(f"Error: {str(e)}")
else:
    st.info("👆 Upload PDF(s) to start. You can select multiple files at once.")